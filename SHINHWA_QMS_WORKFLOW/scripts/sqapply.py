#!/usr/bin/env python3
"""담당자가 작성한 'SQ 양식 작성 요청서'(담당자별 시트)를 SQ 양식 워크북에 반영한다.

사용: python3 scripts/sqapply.py <작성된 요청서.xlsx> [--apply] [--out-dir 폴더]
- 기본은 미리보기(저장 안 함). --apply 일 때만 `03_EDIT/수정완료` 에 새 `_수정본N+1` 로 저장한다(기존 파일은 수정하지 않음).
- 반영 후 `python3 scripts/qms_workflow.py recheck` 로 승인대기에 올리고, 승인·배포는 PC 에서 `python 일괄승인.py`.
- 담당자가 쓰지 않은 값은 만들지 않는다. '점검·검사 항목' 이 비어 있는 행은 건너뛴다.

반영 규칙
- 기록양식(8행 머리줄 `일자|품번/LOT|점검·검사 항목|기준|…`): 9행부터 항목을 한 줄에 하나씩 C열, 기준을 D열.
  기준 줄 수가 항목 줄 수와 같으면 줄별로 대응, 다르면 기준 전체를 D9 에 넣고 보고서에 표시. 주기·작성 시점은 7행.
- 기준서(13행 머리줄 `항목|기준(값)|확인 방법|주기|담당`): 14~21행(최대 8행)에 항목(A)·기준·판정 방법(C)·주기(G).
- 문서상태·제정일·개정일·결재란·보존기간은 바꾸지 않는다.
"""
import sys, re, csv, datetime, glob
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
EDIT = ROOT / "03_EDIT" / "수정완료"
REPORTS = ROOT / "12_OUTPUT" / "REPORTS"
HDR = {"양식번호": None, "워크북": None, "점검·검사 항목": None, "기준·판정 방법": None, "주기·작성 시점": None}


def lines(v):
    return [x.strip() for x in str(v or "").replace("\r", "").split("\n") if x.strip()]


def find_cols(ws):
    for r in range(1, 12):
        vals = [str(c.value or "").replace("\n", "") for c in ws[r]]
        if "양식번호" in vals:
            col = {}
            for i, v in enumerate(vals):
                for k in HDR:
                    if v.startswith(k.replace("\n", "")):
                        col[k] = i
            return r, col
    return None, None


def latest_workbook(name_prefix):
    """요청서 '워크북' 열 값(파일명 앞부분) → 수정완료 의 최신 수정본 파일."""
    base = name_prefix.replace(".xlsx", "")
    base = re.sub(r"_수정본\d*$", "", base)
    cands = [p for p in EDIT.glob(base + "_수정본*.xlsx")]
    if not cands:
        return None, 0
    def num(p):
        m = re.search(r"_수정본(\d*)\.xlsx$", p.name)
        return int(m.group(1)) if m and m.group(1) else 1
    p = max(cands, key=num)
    return p, num(p)


def apply_row(ws, items, stds, cycle):
    """시트에 항목 반영. 반환: 메모 리스트(경고)"""
    notes = []
    kind = "기준서" if str(ws["A13"].value or "").strip() == "항목" else "기록양식"
    if kind == "기록양식":
        first, last = 9, 20
        cap = last - first + 1
        if len(items) > cap:
            notes.append(f"항목 {len(items)}개 중 {cap}개만 반영(초과 {len(items) - cap}개는 작성칸 부족)")
            items = items[:cap]
        aligned = len(stds) == len(items)
        for i, it in enumerate(items):
            ws.cell(first + i, 3).value = it
            if aligned:
                ws.cell(first + i, 4).value = stds[i]
        if stds and not aligned:
            ws.cell(first, 4).value = "\n".join(stds)
            notes.append("기준 줄 수가 항목 수와 달라 기준 전체를 첫 행 D열에 넣음(담당자 확인 필요)")
        if cycle:
            ws["A7"].value = "작성 주기·시점: " + " / ".join(lines(cycle))
    else:
        first, last = 14, 21
        cap = last - first + 1
        if len(items) > cap:
            notes.append(f"항목 {len(items)}개 중 {cap}개만 반영(초과 {len(items) - cap}개는 표 행 부족)")
            items = items[:cap]
        aligned = len(stds) == len(items)
        for i, it in enumerate(items):
            ws.cell(first + i, 1).value = it
            if aligned:
                ws.cell(first + i, 3).value = stds[i]
            if cycle:
                ws.cell(first + i, 7).value = " / ".join(lines(cycle))[:60]
        if stds and not aligned:
            ws.cell(first, 3).value = "\n".join(stds)
            notes.append("기준 줄 수가 항목 수와 달라 기준 전체를 첫 행 C열에 넣음(담당자 확인 필요)")
    return notes


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if not args:
        sys.exit(__doc__)
    apply = "--apply" in sys.argv
    out_dir = EDIT
    if "--out-dir" in sys.argv:
        out_dir = Path(sys.argv[sys.argv.index("--out-dir") + 1])
    req = openpyxl.load_workbook(args[0])
    todo = {}      # 워크북 → [(양식번호, items, stds, cycle, 담당자시트)]
    skipped = []
    for ws in req.worksheets:
        if ws.title.startswith("00_") or ws.title.startswith("01_"):
            continue
        hr, col = find_cols(ws)
        if not col or len(col) < 5:
            continue
        for row in ws.iter_rows(min_row=hr + 1):
            fno = row[col["양식번호"]].value
            if not fno:
                continue
            items = lines(row[col["점검·검사 항목"]].value)
            if not items:
                skipped.append((ws.title, fno, "점검·검사 항목 미작성"))
                continue
            wbname = str(row[col["워크북"]].value or "")
            todo.setdefault(wbname, []).append(
                (str(fno).strip(), items, lines(row[col["기준·판정 방법"]].value), row[col["주기·작성 시점"]].value, ws.title))
    report = [f"# SQ 요청서 반영 {'결과' if apply else '미리보기'} ({datetime.datetime.now():%Y-%m-%d %H:%M})", "",
              f"- 입력: {Path(args[0]).name}", ""]
    saved = []
    for wbname, rows in sorted(todo.items()):
        src, n = latest_workbook(wbname)
        if not src:
            report.append(f"- ✗ 워크북 `{wbname}` 을(를) 03_EDIT/수정완료 에서 찾지 못함 → 건너뜀")
            continue
        wb = openpyxl.load_workbook(src)
        for fno, items, stds, cycle, owner in rows:
            if fno not in wb.sheetnames:
                report.append(f"- ✗ {fno}: 시트 없음({src.name}) → 건너뜀")
                continue
            notes = apply_row(wb[fno], items, stds, cycle)
            report.append(f"- ✓ {fno} ({owner}): 항목 {len(items)}개 반영" + ("".join(f" / ⚠ {x}" for x in notes)))
        newname = re.sub(r"_수정본\d*\.xlsx$", f"_수정본{n + 1}.xlsx", src.name)
        dst = out_dir / newname
        if apply:
            out_dir.mkdir(parents=True, exist_ok=True)
            wb.save(dst)
            saved.append(dst)
            report.append(f"  → 저장: {dst.relative_to(ROOT) if dst.is_relative_to(ROOT) else dst}")
        else:
            report.append(f"  → (미리보기) 저장 예정: {newname}")
    seen = set(); uniq = []
    for t, f, why in skipped:
        if f not in seen:
            seen.add(f); uniq.append((t, f, why))
    report.append(f"\n- 건너뜀(점검·검사 항목 미작성) {len(uniq)}종: " + ", ".join(f for _, f, _ in uniq[:12]) + (" …" if len(uniq) > 12 else ""))
    if apply:
        REPORTS.mkdir(parents=True, exist_ok=True)
        rp = REPORTS / f"sqapply_{datetime.datetime.now():%Y%m%d_%H%M%S}.md"
        rp.write_text("\n".join(report), encoding="utf-8")
        report.append(f"\n보고서: {rp.relative_to(ROOT)}")
        report.append("다음: python3 scripts/qms_workflow.py recheck → 승인대기 확인 → PC 에서 python 일괄승인.py")
    print("\n".join(report))


if __name__ == "__main__":
    main()
