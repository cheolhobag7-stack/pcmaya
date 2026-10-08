#!/usr/bin/env python3
"""SQ 담당자 회신 취합: 회신 받은 확인요청표(xlsx)를 읽어 회신 현황 체크리스트와 반영 후보를 만든다.

사용법 (SHINHWA_QMS_WORKFLOW/99_산출물/회신취합 에서):
  python3 회신취합.py 회신수집            # 미리보기(파일을 만들지 않음)
  python3 회신취합.py 회신수집 --apply    # 체크리스트 사본 + 반영 후보 보고서 생성

- 회신 파일: 담당자가 '확인 결과'를 채운 `첨부_<담당자>_확인요청표.xlsx` 를 회신수집/ 에 넣는다.
- 원본 체크리스트는 수정하지 않는다(새 파일 `…_회신반영_YYYYMMDD[_vN].xlsx`).
- FM Master·증빙 매핑표·문서 상태는 바꾸지 않는다. 반영 후보만 보고서에 낸다(사람이 확정).
"""
import sys, re, glob, datetime, os, csv
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)  # 99_산출물
CHECKLIST = os.path.join(ROOT, "SQ_회신현황_체크리스트_20261008.xlsx")
MAPCSV = os.path.join(ROOT, "SQ_증빙매핑표_전체6분야_20261008_v2.csv")
CHK4 = ["서명본 일치", "배포 수령 사실", "SQ 요구 충족", "실제 기록 보유"]
TODAY = datetime.date.today()


def unique(path):
    if not os.path.exists(path):
        return path
    b, e = os.path.splitext(path)
    n = 2
    while os.path.exists(f"{b}_v{n}{e}"):
        n += 1
    return f"{b}_v{n}{e}"


def read_reply(path):
    """회신 파일 → {SQ NO: {'type','result','memo','who','date'}} , 오류목록"""
    wb = openpyxl.load_workbook(path)
    out, errs = {}, []
    for ws in wb.worksheets:
        if ws.title == "안내":
            continue
        if ws.title == "보유 확인":
            for r in ws.iter_rows(min_row=4, values_only=True):
                if not r[0]:
                    continue
                ans = [r[8], r[9], r[10], r[11]]
                out[int(r[0])] = dict(type="보유 확인", result=" / ".join(f"{k}:{a or '-'}" for k, a in zip(CHK4, ans)),
                                      raw=ans, memo=r[12], who=r[13], date=r[14])
        else:
            for r in ws.iter_rows(min_row=4, values_only=True):
                if not r[0] or not isinstance(r[1], int):
                    continue
                out[int(r[1])] = dict(type=r[0], result=r[8], raw=[r[8]], memo=r[9], who=r[10], date=r[11])
    return out, errs


def suggest(typ, res, memo):
    """회신 → 반영 후보(사람이 확정). 확정 전에는 어떤 문서·번호도 바꾸지 않는다."""
    s = str(res or "")
    if not s.strip() or s.strip() == "-":
        return "미응답 — 재요청"
    if typ == "보유 확인":
        bad = [x for x in re.findall(r":([^/]+)", s) if "보완" in x or "못함" in x]
        return "문제 보고 — 상태 재판정 검토(메모 확인)" if bad else "보유 유지(4가지 확인 완료)"
    if typ == "보완 필요":
        if "작성 완료" in s: return "항목·기준 작성 파일 수령 → 새 수정본 recheck → 승인·배포 → 매핑표 '보유' 갱신 후보"
        if "작성 예정" in s: return "작성 예정일 확인(메모) → 기한 관리"
        if "기존 사용 양식" in s: return "기존 양식 파일 수령 → 내용 확인 후 대표 양식에 반영 검토"
        if "대표 양식으로 충족" in s: return "대표 양식 충족 확인 → 매핑표 보완 필요 유지/재판정 검토"
        if "분리" in s: return "통합 해제·별도 양식 검토(결정 CSV·FM Master)"
    if typ == "미비":
        if "작성 예정" in s: return "작성 예정일 확인(메모) → 작성 후 recheck·승인"
        if "기존 사용 양식" in s: return "기존 양식 파일 수령 → 내용 확인"
        if "해당 없음" in s: return "해당 없음 회신 → 해당 SQ 요구 담당부서 재확인"
        if "충족" in s: return "대표 양식 충족 확인(대표 양식 작성 선행)"
        if "분리" in s: return "통합 해제·별도 양식 검토"
    if typ == "판단 보류":
        if "분리" in s: return "통합 해제 후보 → 결정 CSV·FM Master 갱신(사람 확정)"
        if "대표 양식으로 충족" in s: return "통합 유지 확정 후보 → 매핑표 상태 재판정(보완 필요/보유)"
        if "보완 후 충족" in s: return "대표 양식 보완 필요 → 보완 필요로 재분류 후보"
        if "원본 양식 있음" in s or "승인 증빙 있음" in s: return "원본 양식·승인 증빙 수령 → 릴리스 여부 결정"
        if "원본 양식 없음" in s: return "신규 작성 필요 → 미비/보완 필요로 재분류 후보"
        if "현물 확인 완료" in s: return "현물 증빙 확인 → 매핑표에 현물 보유 반영 후보(증빙 위치 메모)"
        if "기록 양식 필요" in s: return "기록 양식 필요 → 비문서 처리 재검토·양식 작성"
        if "해당 없음" in s: return "해당 없음 회신 → 요구 담당부서 재확인"
        if "구성 충족" in s: return "SH-FM-037 포함 구성 충족 확인 → 재판정 후보"
    return "회신 내용 확인 필요(메모)"


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    apply = "--apply" in sys.argv
    folder = os.path.join(HERE, args[0]) if args else os.path.join(HERE, "회신수집")
    files = sorted(f for f in glob.glob(os.path.join(folder, "*.xlsx")) if not os.path.basename(f).startswith("~$"))
    if not files:
        print(f"회신 파일이 없습니다: {folder}"); return 1
    wb = openpyxl.load_workbook(CHECKLIST)
    d = wb["상세"]
    rows = {}
    for r in range(2, d.max_row + 1):
        if d.cell(r, 1).value is not None:
            rows[(int(d.cell(r, 1).value), d.cell(r, 5).value)] = r
    by_no = {}
    for (no, typ), r in rows.items():
        by_no.setdefault(no, []).append(r)
    report, warn, per = [], [], {}
    for f in files:
        reply, _ = read_reply(f)
        who = os.path.basename(f)
        for no, a in reply.items():
            cands = by_no.get(no)
            if not cands:
                warn.append(f"{who}: SQ NO {no} 이(가) 체크리스트에 없음"); continue
            r = cands[0]
            owner = d.cell(r, 8).value
            answered = any(str(x or "").strip() not in ("", "-") for x in a["raw"])
            per.setdefault(owner, []).append(answered)
            sug = suggest(a["type"], a["result"], a["memo"])
            report.append([no, a["type"], owner, a["result"] or "", a["memo"] or "", a["who"] or "", str(a["date"] or ""), sug, who])
            if apply and answered:
                note = f"{a['result']}" + (f" | {a['memo']}" if a["memo"] else "")
                d.cell(r, 13).value = TODAY if not isinstance(a["date"], (datetime.date, datetime.datetime)) else a["date"]
                d.cell(r, 13).number_format = "yyyy-mm-dd"
                d.cell(r, 14).value = note
                d.cell(r, 15).value = a["who"]
    # 회신 상태: 응답한 행 = 회신 완료. 응답한 담당자의 나머지(미응답) 행 = 일부 회신
    total = {}
    for (no, typ), r in rows.items():
        total.setdefault(d.cell(r, 8).value, []).append(r)
    if apply:
        for owner, rs in total.items():
            got = [r for r in rs if d.cell(r, 14).value]
            if not got:
                continue
            for r in rs:
                if r in got:
                    d.cell(r, 12).value = "회신 완료"
                elif d.cell(r, 12).value != "회신 완료":
                    d.cell(r, 12).value = "일부 회신"
    print(f"회신 파일 {len(files)}개 / 읽은 항목 {len(report)}건 / 경고 {len(warn)}건")
    for o, a in sorted(per.items()):
        print(f"  {o}: 응답 {sum(a)}건 / 요청 {len(total.get(o, []))}건")
    for w in warn: print("  [경고]", w)
    if not apply:
        print("미리보기입니다. --apply 로 체크리스트 사본과 보고서를 만듭니다."); return 0
    out_x = unique(os.path.join(ROOT, f"SQ_회신현황_체크리스트_회신반영_{TODAY:%Y%m%d}.xlsx"))
    wb.save(out_x)
    out_c = unique(os.path.join(HERE, f"회신반영후보_{TODAY:%Y%m%d}.csv"))
    with open(out_c, "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh); w.writerow(["SQ NO", "요청 유형", "담당자", "확인 결과", "메모", "확인자", "확인일", "반영 후보(사람 확정)", "회신 파일"]); w.writerows(report)
    print("체크리스트 사본:", out_x); print("반영 후보:", out_c)
    return 0

if __name__ == "__main__":
    sys.exit(main())
