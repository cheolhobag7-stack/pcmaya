from pathlib import Path
import re, csv
from datetime import datetime

ROOT = Path(__file__).resolve().parents[1]
REVIEW = ROOT/"02_REVIEW"
REGISTER_DIR = ROOT/"00_CONFIG"
DOC_RE = re.compile(r"SH-(QM|QP|WI|FM)-\d{3}", re.I)
REV_RE = re.compile(r"Rev\.\d{2}", re.I)

def find_register():
    candidates = []
    for p in REGISTER_DIR.glob("*"):
        if p.is_file() and "문서관리대장" in p.stem and p.suffix.lower() in {".xlsx",".xlsm"}:
            candidates.append(p)
    return sorted(candidates)[-1] if candidates else None

def read_register(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=False, read_only=True)
    docs = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            vals = [str(c.value).strip() for c in row if c.value is not None]
            joined = " | ".join(vals)
            dm = DOC_RE.search(joined)
            if not dm:
                continue
            rm = REV_RE.search(joined)
            docs[dm.group(0).upper()] = {
                "rev": rm.group(0) if rm else "",
                "sheet": ws.title,
                "row_text": joined[:500]
            }
    return docs

def scan_actual():
    actual = {}
    for base_name in ["01_ORIGINAL","03_EDIT","04_APPROVAL","05_FINAL"]:
        base = ROOT/base_name
        if not base.exists():
            continue
        for p in base.rglob("*"):
            if not p.is_file():
                continue
            dm = DOC_RE.search(p.stem)
            if not dm:
                continue
            rm = REV_RE.search(p.stem)
            actual.setdefault(dm.group(0).upper(), []).append({
                "rev": rm.group(0) if rm else "",
                "path": str(p.relative_to(ROOT))
            })
    return actual

def main():
    reg = find_register()
    if not reg:
        print("문서관리대장 파일을 00_CONFIG에 넣어주세요. 파일명에 '문서관리대장' 포함 필요.")
        return
    registered = read_register(reg)
    actual = scan_actual()
    keys = sorted(set(registered) | set(actual))
    rows = []
    for k in keys:
        r = registered.get(k)
        a = actual.get(k, [])
        status = "PASS"
        issues = []
        if not r:
            status = "FAIL"
            issues.append("실제 파일 존재 / 대장 미등록")
        if not a:
            status = "FAIL"
            issues.append("대장 등록 / 실제 파일 미존재")
        if r and a and r["rev"]:
            revs = {x["rev"] for x in a if x["rev"]}
            if r["rev"] not in revs:
                status = "HOLD"
                issues.append(f"대장 Rev({r['rev']})와 실제 파일 Rev 불일치")
        rows.append({
            "document_no": k,
            "register_rev": r["rev"] if r else "",
            "actual_revs": "; ".join(sorted({x["rev"] for x in a if x["rev"]})),
            "actual_paths": "; ".join(x["path"] for x in a),
            "status": status,
            "issues": " | ".join(issues)
        })
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    out = REVIEW/f"qms_register_compare_{ts}.csv"
    with out.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=["document_no","register_rev","actual_revs","actual_paths","status","issues"])
        w.writeheader()
        w.writerows(rows)
    print(out)

if __name__ == "__main__":
    main()
