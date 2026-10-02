from pathlib import Path
import difflib, re, csv
from datetime import datetime

ROOT = Path(__file__).resolve().parents[1]
DIFF_DIR = ROOT / "02_REVIEW" / "DIFF"
DIFF_DIR.mkdir(parents=True, exist_ok=True)

def extract_docx(path):
    from docx import Document
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)

def extract_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=False, read_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"[SHEET:{ws.title}]")
        for row in ws.iter_rows():
            vals = [str(c.value) for c in row if c.value is not None]
            if vals:
                parts.append(" | ".join(vals))
    return "\n".join(parts)

def extract(path):
    ext = path.suffix.lower()
    if ext == ".docx":
        return extract_docx(path)
    if ext in {".xlsx",".xlsm"}:
        return extract_xlsx(path)
    if ext in {".txt",".md"}:
        return path.read_text(encoding="utf-8", errors="ignore")
    return ""

def main():
    drafts = sorted((ROOT/"03_EDIT"/"AUTO_DRAFT").glob("BATCH_*"))
    if not drafts:
        print("AUTO_DRAFT 배치가 없습니다.")
        return
    batch = drafts[-1]
    manifest = batch / "draft_manifest.csv"
    if not manifest.exists():
        print("draft_manifest.csv 없음")
        return

    with manifest.open("r", encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    summary = []
    for r in rows:
        src = ROOT / r["source"]
        dst = ROOT / r["draft"]
        try:
            a = extract(src).splitlines()
            b = extract(dst).splitlines()
            diff = list(difflib.unified_diff(a,b,fromfile=str(src),tofile=str(dst),lineterm=""))
            out = DIFF_DIR / f"{src.stem}_{ts}.diff.txt"
            out.write_text("\n".join(diff) if diff else "변경 없음", encoding="utf-8")
            summary.append((src.name, str(out.relative_to(ROOT)), "CHANGED" if diff else "NO_CHANGE"))
        except Exception as e:
            summary.append((src.name, "", f"ERROR: {e}"))

    sum_path = DIFF_DIR / f"diff_summary_{ts}.csv"
    with sum_path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["file","diff_report","result"])
        w.writerows(summary)
    print(sum_path)

if __name__ == "__main__":
    main()
