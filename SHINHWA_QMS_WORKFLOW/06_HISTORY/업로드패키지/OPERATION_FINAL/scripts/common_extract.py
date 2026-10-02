from pathlib import Path

def extract_text(path: Path):
    ext = path.suffix.lower()
    if ext == ".docx":
        from docx import Document
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs if p.text]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        for section in doc.sections:
            for hf in (section.header, section.footer):
                for p in hf.paragraphs:
                    if p.text:
                        parts.append(p.text)
        return "\n".join(parts)
    if ext in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=False, read_only=True)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"[SHEET:{ws.title}]")
            for row in ws.iter_rows():
                vals = [str(c.value) for c in row if c.value is not None]
                if vals:
                    parts.append(" | ".join(vals))
        return "\n".join(parts)
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        return "\n".join((p.extract_text() or "") for p in reader.pages)
    if ext in {".txt", ".md", ".csv"}:
        return path.read_text(encoding="utf-8", errors="ignore")
    return ""
