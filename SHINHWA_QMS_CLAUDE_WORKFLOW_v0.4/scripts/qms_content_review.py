from pathlib import Path
import re, csv, json
from datetime import datetime
import yaml

ROOT = Path(__file__).resolve().parents[1]
CFG = ROOT / "00_CONFIG" / "qms_rules.yaml"
ORIGINAL = ROOT / "01_ORIGINAL"
EDIT = ROOT / "03_EDIT"
APPROVAL = ROOT / "04_APPROVAL"
FINAL = ROOT / "05_FINAL"
REVIEW = ROOT / "02_REVIEW"

for d in [REVIEW]:
    d.mkdir(exist_ok=True)

with CFG.open("r", encoding="utf-8") as f:
    CONFIG = yaml.safe_load(f)

ALLOWED_EXT = {".docx", ".xlsx", ".xlsm", ".pdf", ".txt", ".md"}
DOC_PATTERN = re.compile(r"(SH-(QM|QP|WI|FM)-\d{3})", re.I)
REV_PATTERN = re.compile(r"Rev\.\d{2}", re.I)
REF_PATTERN = re.compile(r"\b(?:SH-)?(?:QM|QP|WI|FM)-\d{3}\b", re.I)

def extract_docx_text(path):
    try:
        from docx import Document
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs if p.text]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        for section in doc.sections:
            for hf in (section.header, section.footer):
                for p in hf.paragraphs:
                    if p.text:
                        parts.append(p.text)
        return "\n".join(parts), None
    except Exception as e:
        return "", f"DOCX 읽기 실패: {e}"

def extract_xlsx_text(path):
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=False, read_only=True)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"[SHEET:{ws.title}]")
            for row in ws.iter_rows():
                vals = [str(c.value) for c in row if c.value is not None]
                if vals:
                    parts.append(" | ".join(vals))
        return "\n".join(parts), None
    except Exception as e:
        return "", f"XLSX 읽기 실패: {e}"

def extract_pdf_text(path):
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        parts = []
        for idx, page in enumerate(reader.pages, start=1):
            try:
                txt = page.extract_text() or ""
                parts.append(f"[PAGE:{idx}]\n{txt}")
            except Exception:
                parts.append(f"[PAGE:{idx}]")
        text = "\n".join(parts)
        if not text.strip():
            return "", "PDF 텍스트 추출 결과 없음(스캔형 PDF 가능성)"
        return text, None
    except Exception as e:
        return "", f"PDF 읽기 실패: {e}"

def extract_text(path):
    ext = path.suffix.lower()
    if ext == ".docx":
        return extract_docx_text(path)
    if ext in {".xlsx", ".xlsm"}:
        return extract_xlsx_text(path)
    if ext == ".pdf":
        return extract_pdf_text(path)
    if ext in {".txt", ".md"}:
        try:
            return path.read_text(encoding="utf-8", errors="ignore"), None
        except Exception as e:
            return "", f"텍스트 읽기 실패: {e}"
    return "", "지원하지 않는 확장자"

def normalize_ref(ref):
    u = ref.upper()
    return u if u.startswith("SH-") else "SH-" + u

def all_project_docs():
    docs = {}
    for base in [ORIGINAL, EDIT, APPROVAL, FINAL]:
        if not base.exists():
            continue
        for p in base.rglob("*"):
            if not p.is_file():
                continue
            m = DOC_PATTERN.search(p.stem)
            r = REV_PATTERN.search(p.stem)
            if m:
                docno = m.group(1).upper()
                rev = r.group(0) if r else ""
                docs.setdefault(docno, []).append({
                    "path": str(p.relative_to(ROOT)),
                    "rev": rev
                })
    return docs

PROJECT_DOCS = all_project_docs()

def inspect_file(path):
    result = {
        "file": str(path.relative_to(ROOT)),
        "filename_doc_no": "",
        "content_doc_no": "",
        "filename_rev": "",
        "content_rev": "",
        "company": "NOT_CHECKED",
        "customer": "NOT_CHECKED",
        "retention_3y": "NOT_CHECKED",
        "lot_target_1h": "NOT_CHECKED",
        "prohibited_phrase": "PASS",
        "missing_refs": "",
        "references": "",
        "approval_hint": "UNKNOWN",
        "status": "PASS",
        "issues": []
    }

    if path.suffix.lower() not in ALLOWED_EXT:
        result["status"] = "SKIP"
        result["issues"].append("지원하지 않는 확장자")
        return result

    fm = DOC_PATTERN.search(path.stem)
    if fm:
        result["filename_doc_no"] = fm.group(1).upper()
    else:
        result["status"] = "FAIL"
        result["issues"].append("파일명 문서번호 누락")

    rm = REV_PATTERN.search(path.stem)
    if rm:
        result["filename_rev"] = rm.group(0)
    else:
        if result["status"] == "PASS":
            result["status"] = "HOLD"
        result["issues"].append("파일명 Rev 누락")

    text, err = extract_text(path)
    if err:
        if "스캔형 PDF" in err:
            if result["status"] == "PASS":
                result["status"] = "HOLD"
            result["issues"].append(err)
        else:
            result["status"] = "FAIL"
            result["issues"].append(err)
        return result

    refs = sorted(set(normalize_ref(m.group(0)) for m in REF_PATTERN.finditer(text)))
    result["references"] = "; ".join(refs[:100])

    content_candidates = [r for r in refs if re.fullmatch(r"SH-(QM|QP|WI|FM)-\d{3}", r)]
    if content_candidates:
        result["content_doc_no"] = content_candidates[0]

    revs = REV_PATTERN.findall(text)
    if revs:
        result["content_rev"] = revs[0]

    if result["filename_doc_no"] and result["content_doc_no"] and result["filename_doc_no"] != result["content_doc_no"]:
        result["status"] = "FAIL"
        result["issues"].append(f"파일명/본문 문서번호 불일치: {result['filename_doc_no']} vs {result['content_doc_no']}")

    if result["filename_rev"] and result["content_rev"] and result["filename_rev"].lower() != result["content_rev"].lower():
        result["status"] = "FAIL"
        result["issues"].append(f"파일명/본문 Rev 불일치: {result['filename_rev']} vs {result['content_rev']}")

    accepted_company = CONFIG["company"].get("accepted_names", [CONFIG["company"]["name"]])
    result["company"] = "PASS" if any(x in text for x in accepted_company) else "HOLD"
    if result["company"] == "HOLD":
        result["issues"].append("회사명 본문 확인 필요")

    cust_names = CONFIG.get("customer", {}).get("accepted_names", [])
    result["customer"] = "PASS" if any(x.lower() in text.lower() for x in cust_names) else "INFO"

    ret_patterns = CONFIG.get("record_retention", {}).get("accepted_patterns", [])
    result["retention_3y"] = "PASS" if any(x in text for x in ret_patterns) else "INFO"

    lot_patterns = CONFIG.get("lot_traceability", {}).get("accepted_patterns", [])
    result["lot_target_1h"] = "PASS" if any(x in text for x in lot_patterns) else "INFO"

    found_prohibited = [p for p in CONFIG.get("prohibited_phrases", []) if p in text]
    if found_prohibited:
        result["prohibited_phrase"] = "FAIL"
        result["status"] = "FAIL"
        result["issues"].append("삭제 지시 문구 발견: " + " / ".join(found_prohibited))

    legacy_refs = sorted(set(m.group(0).upper() for m in re.finditer(r"(?<!SH-)\b(?:QM|QP|WI|FM)-\d{3}\b", text, re.I)))
    if legacy_refs:
        if result["status"] == "PASS":
            result["status"] = "HOLD"
        result["issues"].append("SH 접두어 미적용 참조 확인: " + ", ".join(legacy_refs[:20]))

    missing = [r for r in refs if r not in PROJECT_DOCS and r != result["filename_doc_no"]]
    if missing:
        result["missing_refs"] = "; ".join(missing[:50])
        result["status"] = "FAIL"
        result["issues"].append("실제 파일 미존재 참조: " + ", ".join(missing[:20]))

    low = text.lower()
    if any(k in low for k in ["승인완료", "approved", "최종승인"]):
        result["approval_hint"] = "APPROVED_HINT"
    elif any(k in low for k in ["승인대기", "검토중", "작성중"]):
        result["approval_hint"] = "PENDING_HINT"

    return result

def main():
    files = [p for p in ORIGINAL.rglob("*") if p.is_file()]
    rows = [inspect_file(p) for p in files]
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    csv_path = REVIEW / f"qms_content_review_{timestamp}.csv"
    fields = ["file","filename_doc_no","content_doc_no","filename_rev","content_rev",
              "company","customer","retention_3y","lot_target_1h","prohibited_phrase",
              "missing_refs","references","approval_hint","status","issues"]
    with csv_path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        for r in rows:
            rr = r.copy()
            rr["issues"] = " | ".join(rr["issues"])
            w.writerow(rr)

    md = REVIEW / f"qms_content_review_{timestamp}.md"
    lines = ["# QMS v0.3 자동점검 결과",""]
    for st in ["PASS","HOLD","FAIL","SKIP"]:
        lines.append(f"- {st}: {sum(r['status']==st for r in rows)}")
    lines += ["", "## HOLD / FAIL"]
    for r in rows:
        if r["status"] in {"HOLD","FAIL"}:
            lines.append(f"- [{r['status']}] {r['file']} :: {' | '.join(r['issues'])}")
    md.write_text("\n".join(lines), encoding="utf-8")
    print(csv_path)
    print(md)

if __name__ == "__main__":
    main()
