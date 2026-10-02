from pathlib import Path
import re
import csv
import json
from datetime import datetime

ROOT = Path(__file__).resolve().parents[1]
ORIGINAL = ROOT / "01_ORIGINAL"
REVIEW = ROOT / "02_REVIEW"
LOG = ROOT / "99_LOG"

REVIEW.mkdir(exist_ok=True)
LOG.mkdir(exist_ok=True)

ALLOWED_EXT = {".docx", ".xlsx", ".xlsm", ".txt", ".md"}
DOC_PATTERN = re.compile(r"(SH-(QM|QP|WI|FM)-\d{3})", re.I)
REV_PATTERN = re.compile(r"Rev\.\d{2}", re.I)
REF_PATTERN = re.compile(r"\b(?:SH-)?(?:QM|QP|WI|FM)-\d{3}\b", re.I)

COMPANY_NAMES = ["(주)신화에이치앤티", "주식회사 신화에이치앤티", "신화에이치앤티"]
CUSTOMER_NAMES = ["한온시스템", "Hanon Systems", "HANON"]
RETENTION_PATTERNS = ["3년", "3 년", "36개월", "36 개월"]
LOT_TIME_PATTERNS = ["1시간", "1 시간", "60분", "60 분"]

PROHIBITED_PHRASES = [
    "FM Master 재매핑 반영",
    "QP 제3권 FM-030~039 폐기/참조금지",
    "문서통제 주의: 승인된 FM Master와 일치",
    "부록 E. Rev.00 ACTIVE Release 승인결과(AP-01~04)",
    "※ 실제 승인자 성명·서명·일자 입력 후 관리본 배포",
]

def extract_docx_text(path: Path):
    try:
        from docx import Document
        doc = Document(path)
        parts = []
        for p in doc.paragraphs:
            if p.text:
                parts.append(p.text)
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        for section in doc.sections:
            for hf in (section.header, section.footer):
                for p in hf.paragraphs:
                    if p.text:
                        parts.append(p.text)
        return "\n".join(parts), None
    except Exception as e:
        return "", f"DOCX 읽기 실패: {e}"

def extract_xlsx_text(path: Path):
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=False, read_only=True)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"[SHEET:{ws.title}]")
            for row in ws.iter_rows():
                vals = []
                for cell in row:
                    if cell.value is not None:
                        vals.append(str(cell.value))
                if vals:
                    parts.append(" | ".join(vals))
        return "\n".join(parts), None
    except Exception as e:
        return "", f"XLSX 읽기 실패: {e}"

def extract_text(path: Path):
    ext = path.suffix.lower()
    if ext == ".docx":
        return extract_docx_text(path)
    if ext in {".xlsx", ".xlsm"}:
        return extract_xlsx_text(path)
    if ext in {".txt", ".md"}:
        try:
            return path.read_text(encoding="utf-8", errors="ignore"), None
        except Exception as e:
            return "", f"텍스트 읽기 실패: {e}"
    return "", "지원하지 않는 확장자"

def normalize_ref(ref):
    ref = ref.upper()
    if ref.startswith("SH-"):
        return ref
    return "SH-" + ref

def inspect_file(path: Path):
    result = {
        "file": str(path.relative_to(ROOT)),
        "filename_doc_no": "",
        "content_doc_no": "",
        "filename_rev": "",
        "content_rev": "",
        "company": "NOT_CHECKED",
        "customer": "NOT_CHECKED",
        "retention_3y": "NOT_CHECKED",
        "lot_target_1h": "NOT_CHECKED",
        "prohibited_phrase": "PASS",
        "references": "",
        "status": "PASS",
        "issues": []
    }

    if path.suffix.lower() not in ALLOWED_EXT:
        result["status"] = "SKIP"
        result["issues"].append("지원하지 않는 확장자")
        return result

    fm = DOC_PATTERN.search(path.stem)
    if fm:
        result["filename_doc_no"] = fm.group(1).upper()
    else:
        result["status"] = "FAIL"
        result["issues"].append("파일명 문서번호 누락")

    rm = REV_PATTERN.search(path.stem)
    if rm:
        result["filename_rev"] = rm.group(0)
    else:
        result["status"] = "HOLD" if result["status"] == "PASS" else result["status"]
        result["issues"].append("파일명 Rev 누락")

    text, err = extract_text(path)
    if err:
        result["status"] = "FAIL"
        result["issues"].append(err)
        return result

    doc_nos = [normalize_ref(x.group(0)) for x in REF_PATTERN.finditer(text)]
    uniq_refs = sorted(set(doc_nos))
    result["references"] = "; ".join(uniq_refs[:50])

    content_doc_candidates = [r for r in uniq_refs if re.fullmatch(r"SH-(QM|QP|WI|FM)-\d{3}", r)]
    if content_doc_candidates:
        result["content_doc_no"] = content_doc_candidates[0]

    revs = REV_PATTERN.findall(text)
    if revs:
        result["content_rev"] = revs[0]

    if result["filename_doc_no"] and result["content_doc_no"] and result["filename_doc_no"] != result["content_doc_no"]:
        result["status"] = "FAIL"
        result["issues"].append(f"파일명/본문 문서번호 불일치: {result['filename_doc_no']} vs {result['content_doc_no']}")

    if result["filename_rev"] and result["content_rev"] and result["filename_rev"].lower() != result["content_rev"].lower():
        result["status"] = "FAIL"
        result["issues"].append(f"파일명/본문 Rev 불일치: {result['filename_rev']} vs {result['content_rev']}")

    result["company"] = "PASS" if any(x in text for x in COMPANY_NAMES) else "HOLD"
    if result["company"] == "HOLD":
        result["issues"].append("회사명 본문 확인 필요")

    result["customer"] = "PASS" if any(x.lower() in text.lower() for x in CUSTOMER_NAMES) else "INFO"
    result["retention_3y"] = "PASS" if any(x in text for x in RETENTION_PATTERNS) else "INFO"
    result["lot_target_1h"] = "PASS" if any(x in text for x in LOT_TIME_PATTERNS) else "INFO"

    found_prohibited = [p for p in PROHIBITED_PHRASES if p in text]
    if found_prohibited:
        result["prohibited_phrase"] = "FAIL"
        result["status"] = "FAIL"
        result["issues"].append("삭제 지시 문구 발견: " + " / ".join(found_prohibited))

    # Legacy no-prefix references are flagged for review.
    legacy_refs = sorted(set(m.group(0).upper() for m in re.finditer(r"(?<!SH-)\b(?:QM|QP|WI|FM)-\d{3}\b", text, re.I)))
    if legacy_refs:
        if result["status"] == "PASS":
            result["status"] = "HOLD"
        result["issues"].append("SH 접두어 미적용 참조 확인: " + ", ".join(legacy_refs[:20]))

    return result

def write_summary(rows, path):
    counts = {k: sum(r["status"] == k for r in rows) for k in ["PASS", "HOLD", "FAIL", "SKIP"]}
    lines = [
        "# QMS 자동점검 요약",
        "",
        f"- PASS: {counts['PASS']}",
        f"- HOLD: {counts['HOLD']}",
        f"- FAIL: {counts['FAIL']}",
        f"- SKIP: {counts['SKIP']}",
        "",
        "## 확인 필요 문서",
    ]
    for r in rows:
        if r["status"] in {"HOLD", "FAIL"}:
            lines.append(f"- [{r['status']}] {r['file']} :: {' | '.join(r['issues'])}")
    path.write_text("\n".join(lines), encoding="utf-8")

def main():
    files = [p for p in ORIGINAL.rglob("*") if p.is_file()]
    rows = [inspect_file(p) for p in files]
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    csv_path = REVIEW / f"qms_content_review_{timestamp}.csv"
    with csv_path.open("w", newline="", encoding="utf-8-sig") as f:
        fields = [
            "file","filename_doc_no","content_doc_no","filename_rev","content_rev",
            "company","customer","retention_3y","lot_target_1h",
            "prohibited_phrase","references","status","issues"
        ]
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        for r in rows:
            rr = r.copy()
            rr["issues"] = " | ".join(rr["issues"])
            writer.writerow(rr)

    summary_path = REVIEW / f"qms_content_review_{timestamp}.md"
    write_summary(rows, summary_path)

    print("[QMS CONTENT REVIEW COMPLETE]")
    for k in ["PASS","HOLD","FAIL","SKIP"]:
        print(f"{k}: {sum(r['status']==k for r in rows)}")
    print(f"CSV: {csv_path}")
    print(f"SUMMARY: {summary_path}")

if __name__ == "__main__":
    main()
